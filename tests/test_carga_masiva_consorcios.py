"""La carga masiva de consorcios y unidades.

La plantilla tiene tres pestañas —Instrucciones, Ejemplo y Carga— y sólo se lee
«Carga»: una fila por unidad, con los datos del consorcio en su primera fila.
La plantilla anterior (hojas «Consorcios» y «Unidades») se sigue aceptando.
"""

import io

import openpyxl
import pytest


@pytest.fixture
def admin(base, app_modulo):
    app_modulo.app.config['TESTING'] = True
    c = app_modulo.app.test_client()
    with c.session_transaction() as s:
        s['user'] = {'sub': 'auth0|admin', 'email': 'admin@test', 'name': 'Admin', 'role': 'admin'}
    return c


def _subir(admin, wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return admin.post('/api/consorcios/carga-masiva', data={'file': (buf, 'c.xlsx')},
                      content_type='multipart/form-data').get_json()


def _plantilla(admin):
    return openpyxl.load_workbook(io.BytesIO(admin.get('/api/consorcios/plantilla').data))


def _escribir(ws, filas):
    enc = next(i for i, r in enumerate(ws.iter_rows(values_only=True), 1) if r and r[0] == 'consorcio*')
    for n, valores in enumerate(filas, 1):
        for c, v in enumerate(valores, 1):
            ws.cell(row=enc + n, column=c, value=v)


def test_tres_pestanas_visibles(admin, base):
    wb = _plantilla(admin)
    assert [w.title for w in wb.worksheets if w.sheet_state == 'visible'] == ['Instrucciones', 'Ejemplo', 'Carga']


def test_la_plantilla_sin_tocar_no_crea_nada(admin, base):
    d = _subir(admin, _plantilla(admin))
    assert (d['consorcios_creados'], d['unidades_creadas'], d['errores']) == (0, 0, [])


def test_una_fila_por_unidad_crea_el_consorcio_y_sus_unidades(admin, base):
    wb = _plantilla(admin)
    _escribir(wb['Carga'], [
        ['Nuevo Edificio', 'Calle 123', '30-1', 'Juan', '11-1', '1A', '1', 'departamento', 60, 3, 'Ana', 'ana@x'],
        ['Nuevo Edificio', None, None, None, None, '1B', '1', 'cochera', '12,5', None, None, None],
    ])
    d = _subir(admin, wb)
    assert (d['consorcios_creados'], d['unidades_creadas'], d['errores']) == (1, 2, [])
    nuevo = next(c for c in base['consorcios'] if c['nombre'] == 'Nuevo Edificio')
    assert (nuevo['direccion'], nuevo['encargado_nombre']) == ('Calle 123', 'Juan')
    ufs = {u['numero']: u for u in base['unidades_funcionales'] if u['consorcio_id'] == nuevo['id']}
    assert ufs['1A']['superficie_m2'] == 60 and ufs['1A']['ambientes'] == 3
    assert ufs['1B']['tipo'] == 'cochera' and ufs['1B']['superficie_m2'] == 12.5


def test_las_unidades_se_suman_a_un_consorcio_existente(admin, base):
    wb = _plantilla(admin)
    _escribir(wb['Carga'], [['Mío', None, None, None, None, '9Z']])
    d = _subir(admin, wb)
    assert (d['consorcios_creados'], d['consorcios_reutilizados'], d['unidades_creadas']) == (0, 1, 1)


def test_subir_dos_veces_no_duplica(admin, base):
    wb = _plantilla(admin)
    _escribir(wb['Carga'], [['Mío', None, None, None, None, '1A']])
    d = _subir(admin, wb)
    assert (d['unidades_creadas'], d['unidades_omitidas']) == (0, 1)


def test_la_plantilla_vieja_se_sigue_aceptando(admin, base):
    wb = openpyxl.Workbook()
    wb.active.title = 'Consorcios'
    wb['Consorcios'].append(['nombre*', 'direccion'])
    wb['Consorcios'].append(['Viejo', 'Calle 1'])
    u = wb.create_sheet('Unidades')
    u.append(['consorcio*', 'numero*', 'piso', 'tipo', 'superficie_m2', 'vecino_nombre', 'vecino_email'])
    u.append(['Viejo', '2C', '2', 'departamento', 50, '', ''])
    d = _subir(admin, wb)
    assert (d['consorcios_creados'], d['unidades_creadas'], d['errores']) == (1, 1, [])
