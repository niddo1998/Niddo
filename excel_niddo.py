"""Los Excel de Niddo, con la marca de Niddo.

Todos los .xlsx que baja la app —exportes y plantillas de carga masiva— salían
con encabezados violeta (#7C3AED), un color que la marca ya no usa. Este módulo
es el único lugar donde se decide cómo se ve una planilla:

- Arriba, "niddo" y el título del archivo, como en el resto de la app.
- Encabezados en verde Niddo con texto crema, filas alternadas en crema claro.
- Importes con formato de pesos, fechas dd/mm/aaaa, columnas con ancho útil,
  la fila de encabezados congelada y con filtro.

La primera fila con datos NO es la 1: están el logo y el título arriba. Quien
lea estos archivos de vuelta (la carga masiva de gastos acepta el Excel
exportado) tiene que buscar el encabezado con `fila_de_encabezado`.
"""

from datetime import date, datetime

TERRACOTA = 'E8734A'
TERRACOTA_HONDO = 'C4502B'
VERDE = '2F6F5E'
VERDE_BOSQUE = '1E4A3E'
AMARILLO = 'F2B705'
CREMA = 'F6EFE7'
CREMA_CLARO = 'FBF6EF'
TINTA = '2A211C'
GRIS = '8A7F75'
BORDE = 'E7DDD2'

FUENTE = 'Nunito Sans'   # la de la marca; Excel la reemplaza si no está instalada

FORMATO_PESOS = '"$" #,##0.00'
FORMATO_FECHA = 'dd/mm/yyyy'
FORMATO_NUMERO = '#,##0.##'

# Filas que ocupan el logo, el título y el subtítulo antes del encabezado.
FILAS_TITULO = 3


def _estilos():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    borde = Side(style='thin', color=BORDE)
    return {
        'logo': Font(name=FUENTE, bold=True, size=18, color=TINTA),
        'logo_o': Font(name=FUENTE, bold=True, size=18, color=TERRACOTA),
        'titulo': Font(name=FUENTE, bold=True, size=13, color=TINTA),
        'sub': Font(name=FUENTE, size=9, color=GRIS),
        'head_font': Font(name=FUENTE, bold=True, size=10, color=CREMA),
        'head_fill': PatternFill('solid', fgColor=VERDE),
        'celda': Font(name=FUENTE, size=10, color=TINTA),
        'total_font': Font(name=FUENTE, bold=True, size=10, color=TINTA),
        'total_fill': PatternFill('solid', fgColor=CREMA),
        'zebra': PatternFill('solid', fgColor=CREMA_CLARO),
        'ejemplo': Font(name=FUENTE, size=10, italic=True, color=GRIS),
        'borde': Border(bottom=borde),
        'centro': Alignment(horizontal='center', vertical='center', wrap_text=True),
        'izq': Alignment(horizontal='left', vertical='center'),
        'der': Alignment(horizontal='right', vertical='center'),
    }


def _es_monto(nombre):
    n = str(nombre or '').lower()
    return any(p in n for p in ('monto', 'total', 'interés', 'interes', 'importe', 'saldo', 'expensa', 'deuda'))


def _es_fecha(nombre):
    n = str(nombre or '').lower()
    return any(p in n for p in ('fecha', 'vencimiento', 'venc.'))


def _a_fecha(valor):
    """Una fecha ISO ('2026-09-05' o con hora) como date, para que Excel la
    muestre dd/mm/aaaa y la ordene como fecha y no como texto."""
    if isinstance(valor, (date, datetime)) or not valor:
        return valor
    txt = str(valor)[:10]
    try:
        return datetime.strptime(txt, '%Y-%m-%d').date()
    except ValueError:
        return valor


def _a_numero(valor):
    if isinstance(valor, (int, float)) or valor in (None, ''):
        return valor
    try:
        return float(str(valor).replace(',', '.'))
    except ValueError:
        return valor


def encabezado_de_marca(ws, titulo, subtitulo=None, columnas=6):
    """Logo y título en las primeras filas de la hoja."""
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    e = _estilos()
    celda = ws.cell(row=1, column=1)
    try:
        celda.value = CellRichText(
            TextBlock(InlineFont(rFont=FUENTE, b=True, sz=18, color=TINTA), 'nidd'),
            TextBlock(InlineFont(rFont=FUENTE, b=True, sz=18, color=TERRACOTA), 'o'))
    except Exception:
        celda.value = 'niddo'
        celda.font = e['logo']
    ws.row_dimensions[1].height = 26
    t = ws.cell(row=2, column=1, value=titulo)
    t.font = e['titulo']
    sub = subtitulo or f'Generado el {date.today().strftime("%d/%m/%Y")}'
    ws.cell(row=3, column=1, value=sub).font = e['sub']
    ws.sheet_view.showGridLines = False


def tabla(ws, headers, rows, fila=None, totales=None, formatos=None, anchos=None):
    """Escribe una tabla con el estilo de Niddo a partir de `fila`.

    `formatos` es {índice_de_columna: 'pesos' | 'fecha' | 'numero' | 'texto'};
    lo que no se indica se deduce del nombre de la columna. `totales` son los
    índices de columna que se suman en una fila final.
    Devuelve la fila del encabezado.
    """
    from openpyxl.utils import get_column_letter
    e = _estilos()
    fila = fila or (FILAS_TITULO + 2)
    formatos = dict(formatos or {})
    for i, h in enumerate(headers):
        if i not in formatos:
            formatos[i] = 'pesos' if _es_monto(h) else 'fecha' if _es_fecha(h) else 'texto'

    for c, h in enumerate(headers, 1):
        cel = ws.cell(row=fila, column=c, value=h)
        cel.font = e['head_font']
        cel.fill = e['head_fill']
        cel.alignment = e['centro']
    ws.row_dimensions[fila].height = 22

    largo = [len(str(h)) for h in headers]
    for r, row in enumerate(rows, fila + 1):
        for c, val in enumerate(row, 1):
            tipo = formatos.get(c - 1)
            if tipo == 'pesos' or tipo == 'numero':
                val = _a_numero(val)
            elif tipo == 'fecha':
                val = _a_fecha(val)
            cel = ws.cell(row=r, column=c, value=val)
            cel.font = e['celda']
            cel.border = e['borde']
            if (r - fila) % 2 == 0:
                cel.fill = e['zebra']
            if tipo == 'pesos':
                cel.number_format = FORMATO_PESOS
                cel.alignment = e['der']
            elif tipo == 'numero':
                cel.number_format = FORMATO_NUMERO
                cel.alignment = e['der']
            elif tipo == 'fecha':
                cel.number_format = FORMATO_FECHA
                cel.alignment = e['izq']
            texto = f'{val:,.2f}' if isinstance(val, float) else str(val if val is not None else '')
            largo[c - 1] = max(largo[c - 1], min(len(texto), 60))

    ultima = fila + len(rows)
    if totales and rows:
        ultima += 1
        ws.cell(row=ultima, column=1, value='Total').font = e['total_font']
        for c in range(1, len(headers) + 1):
            cel = ws.cell(row=ultima, column=c)
            cel.fill = e['total_fill']
            if c - 1 in totales:
                letra = get_column_letter(c)
                cel.value = f'=SUBTOTAL(9,{letra}{fila + 1}:{letra}{ultima - 1})'
                cel.number_format = FORMATO_PESOS if formatos.get(c - 1) == 'pesos' else FORMATO_NUMERO
                cel.alignment = e['der']
                cel.font = e['total_font']

    for c in range(1, len(headers) + 1):
        ancho = (anchos or {}).get(c - 1) or max(12, min(largo[c - 1] + 4, 60))
        ws.column_dimensions[get_column_letter(c)].width = ancho

    ws.freeze_panes = ws.cell(row=fila + 1, column=1)
    if rows:
        ws.auto_filter.ref = f'A{fila}:{get_column_letter(len(headers))}{fila + len(rows)}'
    return fila


def libro_con_tabla(titulo, headers, rows, nombre_hoja, subtitulo=None, totales=None, formatos=None):
    """Un .xlsx de una sola hoja: logo, título y la tabla. Es el de los exportes."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = nombre_hoja[:31]
    encabezado_de_marca(ws, titulo, subtitulo, len(headers))
    tabla(ws, headers, rows, totales=totales, formatos=formatos)
    return wb


def hoja_instrucciones(ws, titulo, bloques):
    """La hoja "Instrucciones" de una plantilla.

    `bloques` es una lista de (tipo, texto) con tipo 'titulo', 'paso', 'texto',
    'campo' (texto con el nombre del campo en negrita: ('campo', ('nombre', 'qué es'))),
    'nota' o 'espacio'.
    """
    from openpyxl.styles import Alignment, Font, PatternFill
    e = _estilos()
    encabezado_de_marca(ws, titulo, 'Cómo completar esta planilla')
    ws.column_dimensions['A'].width = 26
    ws.column_dimensions['B'].width = 95
    fila = FILAS_TITULO + 2
    for tipo, texto in bloques:
        if tipo == 'espacio':
            fila += 1
            continue
        if tipo == 'titulo':
            c = ws.cell(row=fila, column=1, value=texto)
            c.font = Font(name=FUENTE, bold=True, size=11, color=CREMA)
            for col in (1, 2):
                ws.cell(row=fila, column=col).fill = PatternFill('solid', fgColor=VERDE)
            ws.row_dimensions[fila].height = 20
        elif tipo == 'paso':
            num, txt = texto
            n = ws.cell(row=fila, column=1, value=num)
            n.font = Font(name=FUENTE, bold=True, size=12, color=TERRACOTA)
            n.alignment = Alignment(horizontal='right', vertical='top')
            t = ws.cell(row=fila, column=2, value=txt)
            t.font = e['celda']
            t.alignment = Alignment(wrap_text=True, vertical='top')
        elif tipo == 'campo':
            nombre, txt = texto
            ws.cell(row=fila, column=1, value=nombre).font = Font(name=FUENTE, bold=True, size=10, color=VERDE_BOSQUE)
            t = ws.cell(row=fila, column=2, value=txt)
            t.font = e['celda']
            t.alignment = Alignment(wrap_text=True, vertical='top')
            if (fila % 2) == 0:
                for col in (1, 2):
                    ws.cell(row=fila, column=col).fill = e['zebra']
        elif tipo == 'nota':
            t = ws.cell(row=fila, column=2, value=texto)
            t.font = Font(name=FUENTE, size=9, italic=True, color=GRIS)
            t.alignment = Alignment(wrap_text=True, vertical='top')
        else:
            t = ws.cell(row=fila, column=2, value=texto)
            t.font = e['celda']
            t.alignment = Alignment(wrap_text=True, vertical='top')
        fila += 1
    ws.sheet_view.showGridLines = False


def fila_de_encabezado(ws, reconoce, max_filas=12):
    """La fila donde está el encabezado de la tabla.

    Los archivos que genera Niddo tienen logo y título arriba; los que arma
    alguien a mano o los de versiones viejas lo tienen en la fila 1. Se elige
    la primera fila (de las primeras `max_filas`) con más columnas reconocidas
    por `reconoce(valor) -> bool`.
    """
    mejor, mejor_n = 1, 0
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_filas, values_only=True), 1):
        n = sum(1 for v in (row or ()) if v not in (None, '') and reconoce(v))
        if n > mejor_n:
            mejor, mejor_n = i, n
    return mejor
